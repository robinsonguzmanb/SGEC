import streamlit as st
import sys
import os
sys.path.append(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
import datetime
import io
import pandas as pd
try:
    import PyPDF2
except ImportError:
    PyPDF2 = None
try:
    from pptx import Presentation
except ImportError:
    Presentation = None
from db.database_manager import get_reglas_grado, guardar_simulacion_gym, guardar_conocimiento_gym, buscar_conocimiento_gym, vaciar_conocimiento_gym, obtener_stats_conocimiento
from utils.ia_manager import get_master_prompt
import ollama
def extract_text_from_pdf(file_bytes):
    if not PyPDF2:
        return "[Error: PyPDF2 no está instalado. Añade PyPDF2 a requirements.txt]"
    try:
        reader = PyPDF2.PdfReader(io.BytesIO(file_bytes))
        text = ""
        for page in reader.pages:
            extracted = page.extract_text()
            if extracted:
                text += extracted + "\n"
        return text
    except Exception as e:
        return f"[Error al leer PDF: {e}]"
def extract_text_from_excel(file_bytes):
    try:
        xls = pd.ExcelFile(io.BytesIO(file_bytes))
        text = ""
        for sheet_name in xls.sheet_names:
            df = pd.read_excel(xls, sheet_name=sheet_name)
            text += f"--- Hoja: {sheet_name} ---\n"
            text += df.to_string(index=False) + "\n\n"
        return text
    except Exception as e:
        return f"[Error al leer Excel: {e}]"
def extract_text_from_pptx(file_bytes):
    if not Presentation:
        return "[Error: python-pptx no está instalado. Añade python-pptx a requirements.txt]"
    try:
        prs = Presentation(io.BytesIO(file_bytes))
        text = ""
        for i, slide in enumerate(prs.slides):
            text += f"--- Diapositiva {i+1} ---\n"
            for shape in slide.shapes:
                if hasattr(shape, "text") and shape.text:
                    text += shape.text + "\n"
        return text
    except Exception as e:
        return f"[Error al leer PPTX: {e}]"
def process_uploaded_files(uploaded_files):
    contexto_archivos = ""
    for file in uploaded_files:
        contexto_archivos += f"\n=== Contenido del archivo: {file.name} ===\n"
        file_bytes = file.read()
        if file.name.lower().endswith('.pdf'):
            contexto_archivos += extract_text_from_pdf(file_bytes)
        elif file.name.lower().endswith(('.xlsx', '.xls')):
            contexto_archivos += extract_text_from_excel(file_bytes)
        elif file.name.lower().endswith('.pptx'):
            contexto_archivos += extract_text_from_pptx(file_bytes)
        else:
            contexto_archivos += "[Formato no soportado]\n"
    return contexto_archivos
if 'authenticated' not in st.session_state or not st.session_state['authenticated']:
    st.warning("Debes iniciar sesión primero para acceder a esta página.")
    st.stop()
rol = st.session_state.get('rol')
if rol != "Administrador":
    st.error("No tienes permisos para acceder a esta página. Solo Administradores.")
    st.stop()
st.title("🏋️ Gimnasio de Simulación (Admin Training Ground)")
st.markdown("Calibra la Inteligencia Artificial probando *prompts* masivos y ajustando la temperatura del modelo Llama 3.")
secret_key = st.session_state['secret_key']
reglas_grado = get_reglas_grado(secret_key)
saldo_disponible_mock = 5000000.0
st.sidebar.header("Parámetros del Modelo")
temperatura = st.sidebar.slider("Temperatura (0.0 = Determinista, 1.0 = Creativo)", min_value=0.0, max_value=1.0, value=0.0, step=0.1)
st.subheader("Lógica de Cálculo y System Prompt")
logica_renta_default = "Renta Bruta anual = (sueldo base + gratificación + movilización + colación + otros haberes especiales) x (12 + numero de bonos)"
logica_renta_editado = st.text_area("Lógica de Renta Bruta Anual (Instrucción para la IA):", value=logica_renta_default, height=68)

system_prompt_generado = get_master_prompt(reglas_grado, saldo_disponible_mock, logica_renta_editado)
system_prompt_editado = st.text_area("Puedes ajustar el prompt del sistema para esta prueba:", value=system_prompt_generado, height=250)
st.divider()
st.subheader("Simulación de Caso y Análisis de Documentos")

if 'gym_contexto' not in st.session_state:
    st.session_state['gym_contexto'] = ""

with st.expander("🤖 Asistente de Redacción de Contexto con IA (Ollama)"):
    st.markdown("¿Necesitas ayuda para crear un contexto o caso ficticio? Pídele a la IA que lo redacte por ti.")
    idea_redaccion = st.text_input("Describe la idea del caso (Ej: 'Un colaborador pide aumento por haber terminado un diplomado')")
    if st.button("Redactar Contexto Automáticamente"):
        if idea_redaccion:
            with st.spinner("Redactando caso..."):
                import ollama
                prompt_draft = f"Actúa como un experto en Recursos Humanos. Escribe un contexto ficticio para un caso de simulación basado en esto: '{idea_redaccion}'. Debe ser profesional, detallado pero conciso, e incluir datos ficticios si es necesario (ej: Sueldo, Grado, CECO). No uses saludos, entrega solo el texto del caso."
                messages_draft = [{"role": "user", "content": prompt_draft}]
                try:
                    response_draft = ollama.chat(model='llama3', messages=messages_draft, options={'temperature': temperatura})
                    st.session_state['gym_contexto'] = response_draft['message']['content']
                    st.rerun()
                except Exception as e:
                    st.error(f"Error al conectar con Ollama: {e}")

contexto_editado = st.text_area(
    "Contexto Básico del Caso (Editable):", 
    value=st.session_state['gym_contexto'],
    height=200
)
tab_archivos, tab_texto = st.tabs(["📂 Subir Archivos", "📝 Pegar Texto / Tablas"])

with tab_archivos:
    archivos_subidos = st.file_uploader("Subir documentos para el Cerebro del Gym (PDF, Excel, PPTX)", type=['pdf', 'xlsx', 'xls', 'pptx'], accept_multiple_files=True)
    
    if st.button("📥 Procesar Archivos al Cerebro", type="secondary"):
        if archivos_subidos:
            progress_bar = st.progress(0)
            total_archivos = len(archivos_subidos)
            chunks_guardados = 0
            
            for i, file in enumerate(archivos_subidos):
                file_bytes = file.read()
                texto_extraido = ""
                if file.name.lower().endswith('.pdf'):
                    texto_extraido = extract_text_from_pdf(file_bytes)
                elif file.name.lower().endswith(('.xlsx', '.xls')):
                    texto_extraido = extract_text_from_excel(file_bytes)
                elif file.name.lower().endswith('.pptx'):
                    texto_extraido = extract_text_from_pptx(file_bytes)
                    
                if texto_extraido and not texto_extraido.startswith("[Error"):
                    chunks = guardar_conocimiento_gym(secret_key, file.name, texto_extraido)
                    chunks_guardados += chunks
                    
                progress_bar.progress((i + 1) / total_archivos)
                
            st.success(f"¡Procesamiento completo! Se fragmentaron e ingirieron {chunks_guardados} párrafos en el Cerebro del Gym.")
        else:
            st.warning("Sube al menos un documento primero.")

with tab_texto:
    st.markdown("Puedes pegar directamente políticas, fragmentos de correos, o incluso copiar y pegar tablas de Excel. El sistema extraerá la información para guardarla.")
    titulo_texto = st.text_input("Título o nombre del conocimiento (Ej: 'Tabla de Bandas Salariales 2026')")
    texto_manual = st.text_area("Pega aquí el texto o tabla:", height=200)
    
    if st.button("📥 Guardar Texto en el Cerebro", type="secondary"):
        if texto_manual.strip() and titulo_texto.strip():
            chunks = guardar_conocimiento_gym(secret_key, titulo_texto.strip(), texto_manual)
            st.success(f"¡Conocimiento guardado! Se fragmentó en {chunks} párrafos dentro del Cerebro.")
        else:
            st.warning("Debes ingresar un título y el contenido del texto para guardar.")

# Mostrar estado del Cerebro
stats = obtener_stats_conocimiento(secret_key)
st.info(f"🧠 Estado del Cerebro Local: {stats['docs']} documentos procesados, {stats['chunks']} fragmentos de conocimiento acumulados.")
if stats['chunks'] > 0:
    if st.button("🗑️ Vaciar Cerebro (Reset)", type="secondary"):
        vaciar_conocimiento_gym(secret_key)
        st.success("Cerebro vaciado correctamente.")
        st.rerun()

st.divider()
st.subheader("Simulación de Caso")
prompt_prueba = st.text_area("Ingresa el Prompt de Prueba (Descripción de Tareas / Pregunta sobre políticas):", height=150, placeholder="Ej: Se le asignará la dirección de la estrategia financiera regional...")

if st.button("▶️ Ejecutar Simulación IA", type="primary"):
    if not prompt_prueba:
        st.warning("Debes ingresar un prompt de prueba para simular.")
    else:
        with st.spinner("Buscando en el cerebro e interactuando con Llama 3..."):
            
            contexto_extra = ""
            if stats['chunks'] > 0:
                st.info("🔎 Recuperando conocimiento relevante del Cerebro...")
                resultados_rag = buscar_conocimiento_gym(prompt_prueba, limit=5)
                if resultados_rag:
                    contexto_extra = "=== POLÍTICAS Y CONOCIMIENTO RECUPERADO DE LA EMPRESA ===\n"
                    for i, r in enumerate(resultados_rag):
                        contexto_extra += f"Fragmento {i+1} (Fuente: {r['doc']}):\n{r['texto']}\n\n"
                    with st.expander("Ver fragmentos inyectados (Debug RAG)"):
                        st.text(contexto_extra)
                
            prompt_final = prompt_prueba
            if contexto_extra:
                prompt_final += f"\n\n{contexto_extra}"
                
            messages = [
                {"role": "system", "content": system_prompt_editado},
                {"role": "system", "content": contexto_editado},
                {"role": "user", "content": prompt_final}
            ]
            
            try:
                response = ollama.chat(model='llama3', messages=messages, options={'temperature': temperatura})
                respuesta_ia = response['message']['content']
                
                st.markdown("### 🤖 Respuesta de la IA:")
                
                if respuesta_ia.strip().startswith("[DISCREPANCIA_DETECTADA]") or respuesta_ia.strip().startswith("[PRESUPUESTO_EXCEDIDO]"):
                    st.error("LA IA BLOQUEÓ LA OPERACIÓN")
                else:
                    st.success("LA IA PERMITIÓ LA OPERACIÓN")
                    
                st.info(respuesta_ia)
                
                fecha_str = datetime.datetime.now().strftime("%Y-%m-%d %H:%M:%S")
                guardar_simulacion_gym(secret_key, prompt_prueba, system_prompt_editado, temperatura, respuesta_ia, fecha_str)
                st.success("💾 Simulación guardada en el log de entrenamiento (Simulaciones_Entrenamiento).")
                
            except Exception as e:
                st.error(f"Error al ejecutar Ollama: {e}")